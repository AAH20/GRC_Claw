"""File parsing integrations for different resume formats."""

from __future__ import annotations

from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

import structlog

from resume_parser.models import FileType

logger = structlog.get_logger(__name__)


class BaseFileParser(ABC):
    """Abstract base class for file parsers."""

    @abstractmethod
    def parse(self, file_path: Path) -> str:
        """Parse a file and extract text content.

        Args:
            file_path: Path to the file.

        Returns:
            str: Extracted text content.
        """
        ...

    @abstractmethod
    def supports(self, file_type: FileType) -> bool:
        """Check if this parser supports the given file type.

        Args:
            file_type: File type to check.

        Returns:
            bool: True if supported.
        """
        ...


class PDFParser(BaseFileParser):
    """Parser for PDF files using pypdf."""

    def parse(self, file_path: Path) -> str:
        """Parse a PDF file and extract text.

        Args:
            file_path: Path to the PDF file.

        Returns:
            str: Extracted text content.

        Raises:
            FileParseError: If the PDF cannot be parsed.
        """
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            text_parts: list[str] = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n".join(text_parts)
        except Exception as e:
            logger.error("pdf_parse_error", path=str(file_path), error=str(e))
            raise FileParseError(f"Failed to parse PDF: {e}") from e

    def supports(self, file_type: FileType) -> bool:
        """Check if this parser supports PDF files.

        Args:
            file_type: File type to check.

        Returns:
            bool: True if file type is PDF.
        """
        return file_type == FileType.PDF


class DOCXParser(BaseFileParser):
    """Parser for DOCX files using python-docx."""

    def parse(self, file_path: Path) -> str:
        """Parse a DOCX file and extract text.

        Args:
            file_path: Path to the DOCX file.

        Returns:
            str: Extracted text content.

        Raises:
            FileParseError: If the DOCX cannot be parsed.
        """
        try:
            from docx import Document

            doc = Document(str(file_path))
            text_parts: list[str] = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)
            return "\n".join(text_parts)
        except Exception as e:
            logger.error("docx_parse_error", path=str(file_path), error=str(e))
            raise FileParseError(f"Failed to parse DOCX: {e}") from e

    def supports(self, file_type: FileType) -> bool:
        """Check if this parser supports DOCX files.

        Args:
            file_type: File type to check.

        Returns:
            bool: True if file type is DOCX.
        """
        return file_type == FileType.DOCX


class TXTParser(BaseFileParser):
    """Parser for plain text files."""

    def parse(self, file_path: Path) -> str:
        """Parse a TXT file and extract text.

        Args:
            file_path: Path to the TXT file.

        Returns:
            str: Extracted text content.

        Raises:
            FileParseError: If the file cannot be read.
        """
        try:
            return file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            # Try with latin-1 as fallback
            try:
                return file_path.read_text(encoding="latin-1")
            except Exception as e:
                logger.error("txt_parse_error", path=str(file_path), error=str(e))
                raise FileParseError(f"Failed to parse TXT: {e}") from e
        except Exception as e:
            logger.error("txt_parse_error", path=str(file_path), error=str(e))
            raise FileParseError(f"Failed to parse TXT: {e}") from e

    def supports(self, file_type: FileType) -> bool:
        """Check if this parser supports TXT files.

        Args:
            file_type: File type to check.

        Returns:
            bool: True if file type is TXT.
        """
        return file_type == FileType.TXT


class FileParserRegistry:
    """Registry of file parsers."""

    def __init__(self) -> None:
        """Initialize the parser registry with default parsers."""
        self._parsers: list[BaseFileParser] = [
            PDFParser(),
            DOCXParser(),
            TXTParser(),
        ]

    def get_parser(self, file_type: FileType) -> BaseFileParser:
        """Get a parser for the given file type.

        Args:
            file_type: File type to get parser for.

        Returns:
            BaseFileParser: Appropriate parser.

        Raises:
            UnsupportedFileTypeError: If no parser supports the file type.
        """
        from resume_parser.utils import UnsupportedFileTypeError

        for parser in self._parsers:
            if parser.supports(file_type):
                return parser
        raise UnsupportedFileTypeError(file_type.value)

    def parse_file(self, file_path: Path, file_type: FileType) -> str:
        """Parse a file using the appropriate parser.

        Args:
            file_path: Path to the file.
            file_type: Type of the file.

        Returns:
            str: Extracted text content.
        """
        parser = self.get_parser(file_type)
        return parser.parse(file_path)

    def register_parser(self, parser: BaseFileParser) -> None:
        """Register a custom file parser.

        Args:
            parser: Parser to register.
        """
        self._parsers.append(parser)


# Global registry instance
_registry = FileParserRegistry()


def get_file_parser_registry() -> FileParserRegistry:
    """Get the global file parser registry.

    Returns:
        FileParserRegistry: Global registry instance.
    """
    return _registry
